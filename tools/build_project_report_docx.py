from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
REPORT_MD = ROOT / "项目汇报报告.md"
ASSET_DIR = ROOT / "report-assets"
OUTPUT = ROOT / "转基因时光机_项目汇报报告.docx"


THEME = {
    "navy": "06121A",
    "blue": "11C5FF",
    "green": "74FFB7",
    "yellow": "FFE873",
    "orange": "FFB15C",
    "muted": "5B6770",
    "light": "F5F8FB",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D7DEE8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        border = borders.find(qn(tag))
        if border is None:
            border = OxmlElement(tag)
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_run_font(run, name: str = "Microsoft YaHei", size: Pt | None = None) -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = size


def set_paragraph_format(paragraph, *, after=6, before=0, line=1.25, align=None) -> None:
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(0.86)
    section.right_margin = Inches(0.86)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    for style_name, size, color in [
        ("Heading 1", 16, THEME["blue"]),
        ("Heading 2", 13, THEME["navy"]),
        ("Heading 3", 11.5, "1F4D78"),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_cover(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, THEME["navy"])
    set_cell_border(cell, THEME["blue"])
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("本科《生命科学导论》项目式作业汇报")
    set_run_font(run, size=Pt(12))
    run.font.color.rgb = RGBColor.from_string(THEME["green"])

    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("《转基因时光机》")
    set_run_font(run, size=Pt(28))
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(THEME["yellow"])

    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("从餐桌到医院的生命科学探索")
    set_run_font(run, size=Pt(16))
    run.font.color.rgb = RGBColor.from_string("FFFFFF")

    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, before=16, after=12, line=1.2)
    run = p.add_run("交互式网页项目总结报告")
    set_run_font(run, size=Pt(13))
    run.font.color.rgb = RGBColor.from_string(THEME["orange"])

    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("React + Vite + TypeScript | 黑板手绘 + 未来实验室 UI | localStorage 学习进度")
    set_run_font(run, size=Pt(10))
    run.font.color.rgb = RGBColor.from_string("DDEAF2")

    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, before=18, after=6, line=1.2)
    run = p.add_run("创作标记：崇新学堂 · 何昱隆")
    set_run_font(run, size=Pt(11))
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(THEME["yellow"])

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("报告围绕选题依据、课程内容设计、交互体验、技术实现、项目价值与后续优化进行总结，并配合项目真实页面截图展示。")
    set_run_font(run, size=Pt(10.5))
    run.font.color.rgb = RGBColor.from_string("38414A")

    info = doc.add_table(rows=4, cols=2)
    info.style = "Table Grid"
    rows = [
        ("项目主题", "转基因的前世今生：食品、植物、动物、基因治疗与基因编辑"),
        ("课程定位", "本科《生命科学导论》项目式学习与课堂展示"),
        ("核心目标", "用证据、风险、监管和伦理边界理解基因技术"),
        ("文档日期", "2026 年 7 月 7 日"),
    ]
    for i, (k, v) in enumerate(rows):
        k_cell, v_cell = info.rows[i].cells
        set_cell_shading(k_cell, "EAF7FB")
        set_cell_border(k_cell)
        set_cell_border(v_cell)
        k_cell.text = k
        v_cell.text = v
        for paragraph in k_cell.paragraphs + v_cell.paragraphs:
            set_paragraph_format(paragraph, after=2, line=1.15)
            for run in paragraph.runs:
                set_run_font(run, size=Pt(9.5))
        k_cell.paragraphs[0].runs[0].font.bold = True

    doc.add_page_break()


def add_summary_page(doc: Document) -> None:
    h = doc.add_heading("报告摘要", level=1)
    set_paragraph_format(h, before=0, after=8, line=1.2)

    summary = [
        "本项目把“转基因”从单一概念扩展为一条完整的生命科学学习路线。学习者从 DNA、基因、蛋白质与性状的基础关系出发，逐步进入转基因食品安全、智慧农业、动物研究、基因治疗和 CRISPR 基因编辑等主题，最后通过综合挑战完成技术评价。",
        "项目选择这一主题，是因为转基因技术既能连接课程中的核心概念，也能触达现实生活中的食品安全、农业生产、疾病治疗与伦理监管问题。相比只讲定义，本项目更强调学生在情境中使用知识：看证据是否充分，看风险能否控制，看监管是否到位，看伦理边界是否清楚。",
        "交互设计上，项目采用“黑板手绘 + 未来实验室 UI”的视觉语言，把学习路径做成基因岛屿地图、技术发展时间线和多个角色化任务。技术实现上，项目使用 React + Vite + TypeScript，无后端、无真实 API，所有教学数据写入 src/data，并用 localStorage 保存学习进度，适合课堂演示和学生自主探索。",
    ]
    for text in summary:
        p = doc.add_paragraph(text)
        set_paragraph_format(p, after=7, line=1.32)
        for run in p.runs:
            set_run_font(run, size=Pt(10.5))

    table = doc.add_table(rows=1, cols=3)
    labels = [
        ("课程价值", "把抽象概念转化为可视化、可互动、可复习的学习体验。"),
        ("现实价值", "联系粮食安全、智慧农业、疾病治疗和生命伦理等真实议题。"),
        ("展示价值", "有完整视觉风格、学习进度、互动任务和最终评价，适合课堂汇报。"),
    ]
    for i, (title, body) in enumerate(labels):
        cell = table.cell(0, i)
        set_cell_shading(cell, "F4F8FB")
        set_cell_border(cell, "BFD8E6")
        p = cell.paragraphs[0]
        run = p.add_run(title)
        set_run_font(run, size=Pt(10.5))
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string("0A6D8E")
        p = cell.add_paragraph(body)
        set_paragraph_format(p, after=2, line=1.2)
        for run in p.runs:
            set_run_font(run, size=Pt(9))

    doc.add_page_break()


def load_font(size: int, bold: bool = False):
    candidates = [
        "C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf",
        "C:/Windows/Fonts/simsun.ttc",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def rounded_rect(draw: ImageDraw.ImageDraw, xy, radius: int, fill, outline, width=3) -> None:
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw: ImageDraw.ImageDraw, start, end, color, width=4) -> None:
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    tip = (x2, y2)
    left = (x2 - ux * 20 + px * 9, y2 - uy * 20 + py * 9)
    right = (x2 - ux * 20 - px * 9, y2 - uy * 20 - py * 9)
    draw.polygon([tip, left, right], fill=color)


def generate_structure_diagram() -> Path:
    out = ASSET_DIR / "02-project-structure.png"
    W, H = 1600, 930
    img = Image.new("RGB", (W, H), "#061019")
    draw = ImageDraw.Draw(img)
    title_font = load_font(46, bold=True)
    label_font = load_font(28, bold=True)
    small_font = load_font(22)
    tiny_font = load_font(18)

    for x in range(0, W, 56):
        draw.line([(x, 0), (x, H)], fill="#102332", width=1)
    for y in range(0, H, 42):
        draw.line([(0, y), (W, y)], fill="#102332", width=1)

    draw.text((70, 55), "项目结构图：从主题入口到综合评价", font=title_font, fill="#FFE873")
    draw.text((75, 125), "一条学习路径，五个核心模块，三个支撑系统，最终落到证据化技术评估。", font=small_font, fill="#CFE9F2")

    boxes = [
        ("首页 Hero", "建立主题氛围\n引导开始探索", (80, 230, 330, 370), "#11C5FF"),
        ("时间线", "传统育种到\nCRISPR 的脉络", (410, 230, 660, 370), "#74FFB7"),
        ("知识地图", "基因岛屿路径\n完成后节点点亮", (740, 230, 990, 370), "#FFE873"),
        ("互动实验室", "五个模块入口\n角色化任务", (1070, 230, 1340, 370), "#FFB15C"),
        ("模块学习页", "是什么 / 怎么做\n有什么用 / 风险争议", (260, 520, 580, 700), "#11C5FF"),
        ("复习中心", "进度、错题、术语\n支持课后复习", (660, 520, 980, 700), "#74FFB7"),
        ("最终挑战", "100 分评估试卷\n综合伦理决策", (1060, 520, 1380, 700), "#FFE873"),
    ]
    for title, body, xy, color in boxes:
        rounded_rect(draw, xy, 18, "#0A202A", color, width=4)
        draw.text((xy[0] + 30, xy[1] + 24), title, font=label_font, fill=color)
        y = xy[1] + 74
        for line in body.split("\n"):
            draw.text((xy[0] + 30, y), line, font=small_font, fill="#F5F8FB")
            y += 34

    arrow(draw, (330, 300), (410, 300), "#11C5FF")
    arrow(draw, (660, 300), (740, 300), "#74FFB7")
    arrow(draw, (990, 300), (1070, 300), "#FFE873")
    arrow(draw, (1205, 370), (1220, 520), "#FFB15C")
    arrow(draw, (1070, 330), (580, 560), "#11C5FF")
    arrow(draw, (580, 610), (660, 610), "#74FFB7")
    arrow(draw, (980, 610), (1060, 610), "#FFE873")

    modules = ["食品安全", "智慧农业", "动物伦理", "基因治疗", "基因编辑"]
    x = 180
    for i, module in enumerate(modules, start=1):
        bx = x + (i - 1) * 260
        rounded_rect(draw, (bx, 790, bx + 205, 850), 12, "#081821", "#315E6D", width=2)
        draw.text((bx + 22, 805), f"{i}. {module}", font=tiny_font, fill="#D8F4FF")

    img.save(out, quality=95)
    return out


def add_image_with_caption(doc: Document, image_path: Path, caption: str) -> None:
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(6.25))

    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(cap, after=10, line=1.1)
    for run in cap.runs:
        set_run_font(run, size=Pt(8.5))
        run.font.color.rgb = RGBColor.from_string("5B6770")


def add_figure_bundle(doc: Document, marker_id: str) -> None:
    structure = generate_structure_diagram()
    bundles: dict[str, list[tuple[Path, str]]] = {
        "1": [
            (ASSET_DIR / "01-home.png", "图 1  项目首页主视觉：DNA 线条动画、黑板网格背景与创作标记共同建立“基因技术探索实验室”的氛围。"),
        ],
        "2": [
            (structure, "图 2  项目结构图：从首页、时间线、知识地图进入五个学习模块，并通过复习中心和最终挑战完成闭环。"),
            (ASSET_DIR / "03-timeline.png", "图 3  技术发展时间线：将传统育种、重组 DNA、转基因作物、基因治疗和 CRISPR 放在同一历史脉络中理解。"),
        ],
        "3": [
            (ASSET_DIR / "02-knowledge-map.png", "图 4  知识地图：用“基因岛屿地图”呈现学习路线，模块完成后节点点亮，增强路径感和成就感。"),
        ],
        "4": [
            (ASSET_DIR / "04-interactive-lab.png", "图 5  互动实验室：五个模块以探索舱形式呈现，学生可从任意主题进入学习。"),
        ],
        "5": [
            (ASSET_DIR / "06-module-crispr.png", "图 6  基因编辑模块：用虚构 DNA 片段和伦理讨论提示解释 CRISPR 的概念边界。"),
        ],
        "6": [
            (ASSET_DIR / "05-module-food-safety.png", "图 7  模块学习页：采用分段学习、任务面板和概念卡片，避免大段文字堆叠。"),
        ],
        "7": [
            (ASSET_DIR / "07-final-challenge.png", "图 8  最终挑战：升级为 100 分结课小测，覆盖概念迁移、证据判断、框架匹配和案例分析。"),
            (ASSET_DIR / "08-review-center.png", "图 9  复习中心：汇总学习进度、错题解析和关键词，支持课前展示与课后复习。"),
        ],
        "8": [
            (ASSET_DIR / "09-glossary.png", "图 10  术语抽屉：固定入口帮助学生随时查阅 DNA、GMO、CRISPR、脱靶、伦理审查等概念。"),
        ],
    }
    for image_path, caption in bundles.get(marker_id, []):
        add_image_with_caption(doc, image_path, caption)


def add_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    set_paragraph_format(p, after=7, line=1.32, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    for run in p.runs:
        set_run_font(run, size=Pt(10.5))


def add_list_paragraph(doc: Document, text: str, numbered: bool = False) -> None:
    style = "List Number" if numbered else "List Bullet"
    p = doc.add_paragraph(text, style=style)
    set_paragraph_format(p, after=4, line=1.22)
    for run in p.runs:
        set_run_font(run, size=Pt(10))


def build_docx() -> None:
    ASSET_DIR.mkdir(exist_ok=True)
    md = REPORT_MD.read_text(encoding="utf-8")

    doc = Document()
    style_document(doc)
    props = doc.core_properties
    props.title = "《转基因时光机》项目汇报报告"
    props.subject = "本科生命科学导论交互式课程网站项目总结"
    props.author = "崇新学堂 · 何昱隆"
    props.keywords = "转基因, 基因编辑, 基因治疗, 智慧农业, 生命科学导论, React, Vite, TypeScript"

    add_cover(doc)
    add_summary_page(doc)

    marker_re = re.compile(r"^【配图位置\s*(\d+)：(.+?)】")

    skip_next_suggestion = False
    for raw_line in md.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        marker = marker_re.match(line)
        if marker:
            add_figure_bundle(doc, marker.group(1))
            skip_next_suggestion = True
            continue

        if skip_next_suggestion and line.startswith("建议"):
            skip_next_suggestion = False
            continue
        skip_next_suggestion = False

        if line.startswith("# "):
            continue
        if line.startswith("## "):
            heading = doc.add_heading(line[3:], level=1)
            set_paragraph_format(heading, before=12, after=8, line=1.15)
            continue
        if line.startswith("### "):
            heading = doc.add_heading(line[4:], level=2)
            set_paragraph_format(heading, before=8, after=5, line=1.15)
            continue

        number_match = re.match(r"^\d+\.\s+(.+)$", line)
        if number_match:
            add_list_paragraph(doc, number_match.group(1), numbered=True)
            continue
        if line.startswith("- "):
            add_list_paragraph(doc, line[2:], numbered=False)
            continue

        add_paragraph(doc, line)

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    h = doc.add_heading("附：课堂汇报建议", level=1)
    set_paragraph_format(h, before=0, after=8, line=1.2)
    suggestions = [
        "汇报开场可以先说明为什么选择转基因主题：它既是课程核心概念的综合窗口，也是现实社会中容易被误解、需要证据化判断的公共议题。",
        "展示网页时建议按照“首页主题氛围—知识地图—一个代表性模块—最终挑战—复习中心”的顺序演示，时间有限时可优先展示食品安全模块和最终挑战。",
        "讲项目价值时，不要只强调页面好看，而要突出它训练学生从证据、风险、监管和伦理边界四个角度分析生命科学技术。",
        "讲项目边界时，需要明确说明网页不提供真实实验参数、真实操作流程或真实基因编辑指导，所有 DNA 序列均为虚构示例，项目定位是科普学习与课程展示。",
    ]
    for item in suggestions:
        add_list_paragraph(doc, item, numbered=False)

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_docx()
